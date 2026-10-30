from __future__ import annotations

import csv
import os
import re
from pathlib import Path

import pytest
import yaml
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
PILOT_PATH = ROOT / "data/staging/pilot_accounts.csv"
EXPECTED_PATH = ROOT / "tests/fixtures/pilot_accounts_expected.csv"
SCHEMA_PATH = ROOT / "data/contracts/organization_schema.yaml"
WORKBOOK_PATH = Path(
    os.environ.get("NETWORKING_WORKBOOK", r"C:\Users\Lucas\Desktop\Networking_estagio_v2.xlsx")
)


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open(encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


@pytest.fixture(scope="module")
def pilot_rows() -> list[dict[str, str]]:
    return read_csv(PILOT_PATH)


def test_exactly_twelve_unique_accounts(pilot_rows: list[dict[str, str]]) -> None:
    assert len(pilot_rows) == 12
    assert len({row["account_id"] for row in pilot_rows}) == 12
    assert len({row["canonical_name"] for row in pilot_rows}) == 12


def test_expected_fixture_matches_stable_identity(pilot_rows: list[dict[str, str]]) -> None:
    expected_rows = read_csv(EXPECTED_PATH)
    expected_columns = list(expected_rows[0])
    actual_projection = [{column: row[column] for column in expected_columns} for row in pilot_rows]
    assert actual_projection == expected_rows


def test_required_fields_and_source_locators(pilot_rows: list[dict[str, str]]) -> None:
    required = {
        "account_id",
        "canonical_name",
        "group_name",
        "institution_type",
        "strategy_profile",
        "analysis_unit",
        "source_sheet",
        "source_row",
        "source_locator",
        "raw_aliases",
        "source_workbook_sha256",
    }
    forbidden_placeholders = {"não encontrado", "nao encontrado", "n.a.", "n/a"}
    for row in pilot_rows:
        assert all(row[field].strip() for field in required)
        assert re.fullmatch(r"org-[a-z0-9-]+", row["account_id"])
        assert re.fullmatch(r"[0-9a-f]{64}", row["source_workbook_sha256"])
        assert not any(value.strip().lower() in forbidden_placeholders for value in row.values())
        assert row["evidence_class"] == "source_stated"
        assert row["human_approval_required"] == "true"
        assert row["verification_locator"].startswith("'Base Completa E'!")
        if row["as_of_date"]:
            assert re.fullmatch(r"\d{4}-\d{2}-\d{2}", row["as_of_date"])


def test_bank_accounts_require_desk_level_mapping(pilot_rows: list[dict[str, str]]) -> None:
    bank_rows = [row for row in pilot_rows if row["institution_type"] == "banco"]
    assert {row["canonical_name"] for row in bank_rows} == {"BTG Pactual", "Itaú BBA"}
    for row in bank_rows:
        assert row["analysis_unit"] == "mesa/área"
        assert row["business_unit_required"] == "true"
        assert "desk_level_mapping_required" in row["data_quality_flags"]


def test_schema_covers_every_csv_column(pilot_rows: list[dict[str, str]]) -> None:
    schema = yaml.safe_load(SCHEMA_PATH.read_text(encoding="utf-8"))
    assert schema["schema_version"] == "1.0.0"
    assert set(schema["fields"]) == set(pilot_rows[0])
    for field in schema["fields"].values():
        assert {"type", "nullable", "validation", "unit", "origin", "date_semantics"} <= set(field)


def test_source_locators_resolve_to_selected_names(pilot_rows: list[dict[str, str]]) -> None:
    workbook = load_workbook(WORKBOOK_PATH, read_only=True, data_only=True)
    for row in pilot_rows:
        source_sheet = row["source_sheet"]
        source_row = int(row["source_row"])
        source_name = str(workbook[source_sheet].cell(source_row, 1).value)
        assert source_name in row["raw_aliases"].split(" | ")


def test_reconciliation_controls_match_current_workbook() -> None:
    workbook = load_workbook(WORKBOOK_PATH, read_only=False, data_only=False)
    category_sheets = [
        "PE E",
        "Hedge Fund E",
        "Quant E",
        "Equity Research E",
        "Asset E",
        "IB E",
        "ML-DS E",
        "Venture Capital E",
        "Crédito E",
        "Family Office E",
        "Real Estate E",
    ]
    raw_count = sum(
        1
        for sheet_name in category_sheets
        for row in workbook[sheet_name].iter_rows(min_row=2, values_only=True)
        if row[0] and row[1]
    )
    base_count = sum(
        1 for row in workbook["Base Completa E"].iter_rows(min_row=5, values_only=True) if row[1]
    )
    assert raw_count == 396
    assert base_count == 211
    assert workbook["Resumo E"]["B4"].value == "=COUNTA('Base Completa E'!$B$5:$B$214)"
    assert workbook["Progresso"]["B2"].value == "287"
    assert workbook["Base Completa E"]["D215"].value.startswith("https://")
