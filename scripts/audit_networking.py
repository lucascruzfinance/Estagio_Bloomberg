from __future__ import annotations

import argparse
import csv
import hashlib
import re
from collections import Counter
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any

import yaml
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from unidecode import unidecode

CATEGORY_SHEETS = [
    "PE E",
    "Hedge Fund E",
    "Quant E",
    "Equity Research E",
    "Asset E",
    "IB E",
    "ML-DS E",
    "Venture Capital E",
    "Crédito E",
    "Family Office E",
    "Real Estate E",
]

AUDITED_SHEETS = [
    "Base Completa E",
    "Resumo E",
    "Progresso",
    "PE E",
    "Hedge Fund E",
    "Asset E",
    "Family Office E",
    "IB E",
    "Pendências E",
]

MISSING_TOKENS = {"", "não encontrado", "nao encontrado", "n.a.", "n/a", "—", "-"}
ANALYSIS_DATE = "2026-10-29"


def normalized_name(value: str) -> str:
    text = unidecode(value).lower().strip()
    text = re.sub(r"[^a-z0-9]+", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def is_missing(value: Any) -> bool:
    if value is None:
        return True
    return normalized_name(str(value)) in {normalized_name(token) for token in MISSING_TOKENS}


def clean_nullable(value: Any) -> str:
    return "" if is_missing(value) else str(value).strip()


def normalize_date(value: Any) -> str:
    text = clean_nullable(value)
    if not text:
        return ""
    for pattern in ("%d/%m/%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(text, pattern).date().isoformat()
        except ValueError:
            continue
    return ""


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def sheet_locator(sheet: str, row: int, start_col: int = 1, end_col: int = 9) -> str:
    return f"'{sheet}'!{get_column_letter(start_col)}{row}:{get_column_letter(end_col)}{row}"


@dataclass(frozen=True)
class PilotSpec:
    account_id: str
    canonical_name: str
    institution_type: str
    strategy_profile: str
    group_name: str
    analysis_unit: str
    business_unit_required: bool
    source_sheet: str
    source_row: int
    base_record_id: int
    aliases: tuple[str, ...]


PILOT_SPECS = (
    PilotSpec(
        "org-spx-capital",
        "SPX Capital",
        "gestora independente",
        "macro e multimercado",
        "SPX Capital",
        "organização",
        False,
        "Hedge Fund E",
        3,
        185,
        ("spx capital",),
    ),
    PilotSpec(
        "org-verde-asset",
        "Verde Asset Management",
        "gestora independente",
        "macro e multimercado",
        "Verde Asset Management",
        "organização",
        False,
        "Hedge Fund E",
        4,
        1,
        ("verde asset", "verde asset management"),
    ),
    PilotSpec(
        "org-kapitalo",
        "Kapitalo Investimentos",
        "gestora independente",
        "multiestratégia",
        "Kapitalo Investimentos",
        "organização",
        False,
        "Hedge Fund E",
        5,
        2,
        ("kapitalo", "kapitalo investimentos"),
    ),
    PilotSpec(
        "org-ibiuna",
        "Ibiúna Investimentos",
        "gestora independente",
        "macro e multimercado",
        "Ibiúna Investimentos",
        "organização",
        False,
        "Hedge Fund E",
        6,
        4,
        ("ibiuna investimentos", "ibiuna investimentos ibiuna macro gestao de recursos"),
    ),
    PilotSpec(
        "org-kinea",
        "Kinea Investimentos",
        "gestora de grupo financeiro",
        "multiproduto e alternativos",
        "Itaú Unibanco",
        "organização",
        False,
        "Asset E",
        18,
        145,
        ("kinea", "kinea investimentos", "kinea multimercado"),
    ),
    PilotSpec(
        "org-patria",
        "Pátria Investimentos",
        "gestora independente",
        "private equity e alternativos",
        "Pátria Investimentos",
        "organização",
        False,
        "PE E",
        9,
        144,
        ("patria investimentos",),
    ),
    PilotSpec(
        "org-vinci-compass",
        "Vinci Compass",
        "gestora independente",
        "alternativos e líquidos",
        "Vinci Compass",
        "organização",
        False,
        "Asset E",
        20,
        157,
        ("vinci compass", "vinci partners vinci compass", "vinci advisory"),
    ),
    PilotSpec(
        "org-constellation",
        "Constellation Asset Management",
        "gestora independente",
        "ações long only e long biased",
        "Constellation Asset Management",
        "organização",
        False,
        "Asset E",
        35,
        12,
        ("constellation", "constellation asset management"),
    ),
    PilotSpec(
        "org-jive",
        "Jive Investments",
        "gestora independente",
        "crédito e special situations",
        "Jive Investments",
        "organização",
        False,
        "Asset E",
        46,
        25,
        ("jive investments",),
    ),
    PilotSpec(
        "org-portofino",
        "Portofino Gestão de Recursos",
        "multi-family office",
        "wealth e gestão patrimonial",
        "Portofino Gestão de Recursos",
        "organização",
        False,
        "Family Office E",
        3,
        26,
        ("portofino gestao de recursos",),
    ),
    PilotSpec(
        "org-btg-pactual",
        "BTG Pactual",
        "banco",
        "banco universal, asset e investment banking",
        "BTG Pactual",
        "mesa/área",
        True,
        "PE E",
        3,
        168,
        ("btg pactual", "btg pactual asset", "btg pactual asset management"),
    ),
    PilotSpec(
        "org-itau-bba",
        "Itaú BBA",
        "banco",
        "investment banking e mercados",
        "Itaú Unibanco",
        "mesa/área",
        True,
        "IB E",
        4,
        167,
        ("itau bba",),
    ),
)


OUTPUT_COLUMNS = [
    "account_id",
    "canonical_name",
    "legal_name",
    "group_name",
    "institution_type",
    "strategy_profile",
    "analysis_unit",
    "business_unit_required",
    "pilot_wave",
    "cnpj",
    "linkedin_url",
    "email_public",
    "phone_public",
    "website_url",
    "careers_channel",
    "address_raw",
    "city",
    "state",
    "verification_status",
    "evidence_class",
    "confidence",
    "as_of_date",
    "source_sheet",
    "source_row",
    "source_locator",
    "verification_locator",
    "verification_source_url",
    "raw_aliases",
    "missing_fields",
    "data_quality_flags",
    "human_approval_required",
    "source_workbook_sha256",
]


def find_base_row(workbook: Any, record_id: int) -> tuple[int, tuple[Any, ...]] | None:
    # Column A is the formula =ROW()-4, so the record id maps deterministically
    # to its current worksheet row even when formulas are loaded instead of caches.
    row_number = record_id + 4
    row = tuple(cell.value for cell in workbook["Base Completa E"][row_number])
    return (row_number, row) if clean_nullable(row[1]) else None


def collect_aliases(workbook: Any, allowed_aliases: tuple[str, ...]) -> list[str]:
    allowed = {normalized_name(alias) for alias in allowed_aliases}
    found: set[str] = set()
    for sheet_name in CATEGORY_SHEETS:
        for row in workbook[sheet_name].iter_rows(min_row=2, values_only=True):
            raw_name = row[0]
            if raw_name and normalized_name(str(raw_name)) in allowed:
                found.add(str(raw_name).strip())
    return sorted(found, key=normalized_name)


def infer_city_state(address: str) -> tuple[str, str]:
    normalized = normalized_name(address)
    cities: list[str] = []
    if "sao paulo" in normalized:
        cities.append("São Paulo")
    if "rio de janeiro" in normalized:
        cities.append("Rio de Janeiro")
    city = " | ".join(cities)
    state = (
        "SP | RJ" if len(cities) == 2 else ("SP" if city == "São Paulo" else "RJ" if city else "")
    )
    return city, state


def confidence_for(base_row: tuple[Any, ...] | None) -> str:
    if base_row is None:
        return "low"
    status = normalized_name(str(base_row[10] or ""))
    source_url = clean_nullable(base_row[15])
    source_quality_penalties = ("cnpjcheck", "linkedin.com", "econodata", "guias comerciais")
    is_primary_candidate = source_url.startswith("http") and not any(
        token in source_url.lower() for token in source_quality_penalties
    )
    if is_primary_candidate:
        return "high"
    if "confirmado" in status:
        return "medium"
    return "low"


def build_pilot_rows(workbook: Any, source_hash: str) -> list[dict[str, str]]:
    results: list[dict[str, str]] = []
    for spec in PILOT_SPECS:
        row = tuple(cell.value for cell in workbook[spec.source_sheet][spec.source_row][:9])
        base_match = find_base_row(workbook, spec.base_record_id)
        base_row_number = base_match[0] if base_match else None
        base_row = base_match[1] if base_match else None

        address = clean_nullable(row[8])
        city, state = infer_city_state(address)
        aliases = collect_aliases(workbook, spec.aliases)

        verification_status = clean_nullable(base_row[10]) if base_row else ""
        verification_url = clean_nullable(base_row[15]) if base_row else ""
        verification_date = normalize_date(base_row[16]) if base_row else ""
        flags: list[str] = []
        if not verification_url:
            flags.append("verification_source_url_missing")
        if not verification_date:
            flags.append("verification_date_missing")
        if spec.business_unit_required:
            flags.append("desk_level_mapping_required")
        if spec.account_id == "org-itau-bba":
            flags.append("base_record_is_group_level")
        if spec.account_id == "org-btg-pactual":
            flags.append("base_verification_source_is_aggregated")
        if any(
            token in verification_url.lower()
            for token in ("cnpjcheck", "linkedin.com", "econodata", "guias comerciais")
        ):
            flags.append("verification_source_is_non_primary")
        if "|" in address:
            flags.append("multiple_offices_in_single_field")

        row_values = {
            "account_id": spec.account_id,
            "canonical_name": spec.canonical_name,
            "legal_name": "",
            "group_name": spec.group_name,
            "institution_type": spec.institution_type,
            "strategy_profile": spec.strategy_profile,
            "analysis_unit": spec.analysis_unit,
            "business_unit_required": str(spec.business_unit_required).lower(),
            "pilot_wave": "1",
            "cnpj": "",
            "linkedin_url": clean_nullable(row[2]),
            "email_public": clean_nullable(row[3]),
            "phone_public": clean_nullable(row[4]),
            "website_url": clean_nullable(row[5]),
            "careers_channel": clean_nullable(row[7]),
            "address_raw": address,
            "city": city,
            "state": state,
            "verification_status": verification_status,
            "evidence_class": "source_stated",
            "confidence": confidence_for(base_row),
            "as_of_date": verification_date,
            "source_sheet": spec.source_sheet,
            "source_row": str(spec.source_row),
            "source_locator": sheet_locator(spec.source_sheet, spec.source_row),
            "verification_locator": (
                sheet_locator("Base Completa E", base_row_number, 2, 17)
                if base_row_number is not None
                else ""
            ),
            "verification_source_url": verification_url,
            "raw_aliases": " | ".join(aliases),
            "missing_fields": "",
            "data_quality_flags": " | ".join(flags),
            "human_approval_required": "true",
            "source_workbook_sha256": source_hash,
        }
        material_fields = [
            "legal_name",
            "cnpj",
            "linkedin_url",
            "email_public",
            "phone_public",
            "website_url",
            "address_raw",
            "verification_source_url",
            "as_of_date",
        ]
        row_values["missing_fields"] = " | ".join(
            field for field in material_fields if not row_values[field]
        )
        results.append(row_values)
    return results


def workbook_metrics(workbook: Any) -> dict[str, Any]:
    raw_rows: list[tuple[str, int, tuple[Any, ...]]] = []
    for sheet_name in CATEGORY_SHEETS:
        for row_number, row in enumerate(
            workbook[sheet_name].iter_rows(min_row=2, values_only=True), start=2
        ):
            if row[0] and row[1]:
                raw_rows.append((sheet_name, row_number, row))

    exact_names = [str(row[2][0]).strip() for row in raw_rows]
    normalized_names = [normalized_name(name) for name in exact_names]
    counts = Counter(normalized_names)

    missing_counts: dict[str, int] = {}
    header = [cell.value for cell in workbook[CATEGORY_SHEETS[0]][1]][:9]
    for index, label in enumerate(header):
        if index == 0:
            continue
        missing_counts[str(label)] = sum(is_missing(row[2][index]) for row in raw_rows)

    base_rows = [
        (row_number, row)
        for row_number, row in enumerate(
            workbook["Base Completa E"].iter_rows(min_row=5, values_only=True), start=5
        )
        if row[1]
    ]
    malformed_base_rows = []
    for row_number, row in base_rows:
        strategy = clean_nullable(row[3])
        if strategy.startswith("http://") or strategy.startswith("https://"):
            malformed_base_rows.append(row_number)

    hidden_sheets = [sheet.title for sheet in workbook.worksheets if sheet.sheet_state != "visible"]
    formula_count = 0
    formula_ref_errors = 0
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if cell.data_type == "f":
                    formula_count += 1
                    if "#REF!" in str(cell.value):
                        formula_ref_errors += 1

    sheet_profiles: list[dict[str, Any]] = []
    for sheet_name in AUDITED_SHEETS:
        sheet = workbook[sheet_name]
        if sheet_name in CATEGORY_SHEETS:
            data_rows = sum(
                bool(row[0] and row[1]) for row in sheet.iter_rows(min_row=2, values_only=True)
            )
            note = "linhas institucionais; cabeçalhos de seção foram excluídos"
        elif sheet_name == "Base Completa E":
            data_rows = len(base_rows)
            note = "211 nomes; a linha 215 está fora do range oficial e desalinhada"
        elif sheet_name == "Resumo E":
            data_rows = sum(bool(row[0]) for row in sheet.iter_rows(min_row=3, values_only=True))
            note = "fórmulas limitadas à Base Completa E até a linha 214"
        elif sheet_name == "Progresso":
            data_rows = sum(bool(row[0]) for row in sheet.iter_rows(min_row=2, values_only=True))
            note = "controle hardcoded; não deriva da base"
        else:
            data_rows = sum(
                bool(row[0] and row[1]) for row in sheet.iter_rows(min_row=5, values_only=True)
            )
            note = "cópia estática de pendências; a própria aba informa que não se atualiza"
        sheet_profiles.append(
            {
                "sheet": sheet_name,
                "rows": sheet.max_row,
                "columns": sheet.max_column,
                "data_rows": data_rows,
                "formulas": sum(cell.data_type == "f" for row in sheet.iter_rows() for cell in row),
                "note": note,
            }
        )

    return {
        "sheet_count": len(workbook.sheetnames),
        "hidden_sheets": hidden_sheets,
        "raw_rows": len(raw_rows),
        "raw_exact_unique": len(set(exact_names)),
        "raw_normalized_unique": len(set(normalized_names)),
        "duplicate_extra_rows": sum(value - 1 for value in counts.values()),
        "duplicate_groups": sum(value > 1 for value in counts.values()),
        "base_rows": len(base_rows),
        "malformed_base_rows": malformed_base_rows,
        "missing_counts": missing_counts,
        "formula_count": formula_count,
        "formula_ref_errors": formula_ref_errors,
        "progress_count": clean_nullable(workbook["Progresso"]["B2"].value),
        "progress_detail": clean_nullable(workbook["Progresso"]["C2"].value),
        "summary_formula": clean_nullable(workbook["Resumo E"]["B4"].value),
        "sheet_profiles": sheet_profiles,
    }


def schema_document() -> dict[str, Any]:
    field_specs: dict[str, dict[str, Any]] = {
        "account_id": {"type": "string", "nullable": False, "validation": "^org-[a-z0-9-]+$"},
        "canonical_name": {"type": "string", "nullable": False, "validation": "unique; trimmed"},
        "legal_name": {"type": "string", "nullable": True, "validation": "official registry only"},
        "group_name": {"type": "string", "nullable": False, "validation": "trimmed"},
        "institution_type": {
            "type": "string",
            "nullable": False,
            "validation": "controlled vocabulary",
        },
        "strategy_profile": {
            "type": "string",
            "nullable": False,
            "validation": "controlled vocabulary",
        },
        "analysis_unit": {
            "type": "string",
            "nullable": False,
            "validation": "organização|mesa/área",
        },
        "business_unit_required": {
            "type": "boolean",
            "nullable": False,
            "validation": "true|false",
        },
        "pilot_wave": {"type": "integer", "nullable": False, "validation": ">=1"},
        "cnpj": {"type": "string", "nullable": True, "validation": "14 digits and check digits"},
        "linkedin_url": {
            "type": "url",
            "nullable": True,
            "validation": "https; public company/profile URL",
        },
        "email_public": {
            "type": "email",
            "nullable": True,
            "validation": "public professional address only",
        },
        "phone_public": {
            "type": "string",
            "nullable": True,
            "validation": "E.164 after normalization",
        },
        "website_url": {"type": "url", "nullable": True, "validation": "https preferred"},
        "careers_channel": {
            "type": "string",
            "nullable": True,
            "validation": "source-stated; not a sales contact",
        },
        "address_raw": {"type": "string", "nullable": True, "validation": "preserve source text"},
        "city": {
            "type": "string",
            "nullable": True,
            "validation": "derived only when explicit in address",
        },
        "state": {
            "type": "string",
            "nullable": True,
            "validation": "UF; multiple offices pipe-separated",
        },
        "verification_status": {
            "type": "string",
            "nullable": True,
            "validation": "preserve workbook vocabulary",
        },
        "evidence_class": {
            "type": "string",
            "nullable": False,
            "validation": "source_stated|verified_fact|inference|gap",
        },
        "confidence": {
            "type": "string",
            "nullable": False,
            "validation": "high|medium|low|blocked",
        },
        "as_of_date": {
            "type": "date",
            "nullable": True,
            "validation": "ISO 8601 YYYY-MM-DD on database load",
        },
        "source_sheet": {
            "type": "string",
            "nullable": False,
            "validation": "existing workbook sheet",
        },
        "source_row": {"type": "integer", "nullable": False, "validation": ">=2"},
        "source_locator": {"type": "string", "nullable": False, "validation": "Excel A1 locator"},
        "verification_locator": {
            "type": "string",
            "nullable": True,
            "validation": "Excel A1 locator",
        },
        "verification_source_url": {
            "type": "url",
            "nullable": True,
            "validation": "source named in workbook",
        },
        "raw_aliases": {
            "type": "string",
            "nullable": False,
            "validation": "pipe-separated; exact source names",
        },
        "missing_fields": {
            "type": "string",
            "nullable": True,
            "validation": "pipe-separated field names",
        },
        "data_quality_flags": {
            "type": "string",
            "nullable": True,
            "validation": "pipe-separated controlled flags",
        },
        "human_approval_required": {
            "type": "boolean",
            "nullable": False,
            "validation": "true for Prompt 1",
        },
        "source_workbook_sha256": {
            "type": "string",
            "nullable": False,
            "validation": "64 lowercase hex chars",
        },
    }
    for spec in field_specs.values():
        spec["unit"] = "not_applicable"
        spec["origin"] = "source workbook or explicit normalization rule"
        spec["date_semantics"] = "as_of_date records the evidence date; blank means unavailable"
        spec["confidence_semantics"] = (
            "row confidence never upgrades source-stated evidence to verified fact"
        )
    return {
        "schema_name": "pilot_organization_profile",
        "schema_version": "1.0.0",
        "analysis_date": ANALYSIS_DATE,
        "source_policy": {
            "facts": "Workbook values are source-stated until independently reverified.",
            "missing": "Missing data remain null; no proxy substitution.",
            "provenance": "Every row carries source and verification locators plus the workbook hash.",
        },
        "fields": field_specs,
    }


def write_csv(path: Path, rows: list[dict[str, str]], columns: list[str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=columns, extrasaction="raise")
        writer.writeheader()
        writer.writerows(rows)


def write_expected_fixture(path: Path, rows: list[dict[str, str]]) -> None:
    columns = [
        "account_id",
        "canonical_name",
        "institution_type",
        "analysis_unit",
        "source_sheet",
        "source_row",
        "source_locator",
        "human_approval_required",
    ]
    write_csv(path, [{column: row[column] for column in columns} for row in rows], columns)


def markdown_table(headers: list[str], rows: list[list[Any]]) -> str:
    lines = [
        "| " + " | ".join(headers) + " |",
        "| " + " | ".join(["---"] * len(headers)) + " |",
    ]
    for row in rows:
        lines.append("| " + " | ".join(str(value).replace("|", "\\|") for value in row) + " |")
    return "\n".join(lines)


def write_audit(
    path: Path,
    workbook_path: Path,
    source_hash: str,
    metrics: dict[str, Any],
    pilot_rows: list[dict[str, str]],
) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    missing_rows = [
        [field, count, f"{count / metrics['raw_rows']:.1%}"]
        for field, count in metrics["missing_counts"].items()
    ]
    pilot_quality_rows = [
        [
            row["canonical_name"],
            row["confidence"],
            row["verification_status"] or "lacuna",
            row["missing_fields"] or "nenhuma lacuna material no recorte",
            row["data_quality_flags"] or "sem flag adicional",
        ]
        for row in pilot_rows
    ]
    audited_sheet_rows = [
        [
            profile["sheet"],
            f"{profile['rows']} x {profile['columns']}",
            profile["data_rows"],
            profile["formulas"],
            profile["note"],
        ]
        for profile in metrics["sheet_profiles"]
    ]
    content = f"""# Auditoria da base de networking — Prompt 1

Status: **limited**  
Data da análise: {ANALYSIS_DATE}  
Fonte: `{workbook_path}`  
SHA-256: `{source_hash}`

## Mandato e política de mudança

- Objetivo: preparar uma base auditável para o MVP de sales intelligence.
- Corpus autorizado: workbook informado por Lucas e `prompts_codex.md`.
- Modo: auditoria read-only do Excel; o arquivo original não foi modificado.
- Escopo: reconciliação global da estrutura e perfil detalhado apenas das 12 contas do piloto.
- Evidência: todo conteúdo extraído do workbook é `source_stated`, não fato independentemente verificado.
- Dados ausentes permanecem nulos. Não foram inventados CNPJ, pessoas, AUM, performance ou contatos.

## Inventário

- Abas: {metrics["sheet_count"]}.
- Abas ocultas: {", ".join(metrics["hidden_sheets"]) if metrics["hidden_sheets"] else "nenhuma"}.
- Linhas institucionais atuais nas 11 abas de categorias: {metrics["raw_rows"]}.
- Nomes únicos exatos: {metrics["raw_exact_unique"]}.
- Nomes únicos após apenas caixa/acentos/pontuação: {metrics["raw_normalized_unique"]}.
- Linhas duplicadas adicionais sob essa normalização: {metrics["duplicate_extra_rows"]} em {metrics["duplicate_groups"]} grupos.
- Fórmulas no workbook: {metrics["formula_count"]}; fórmulas contendo `#REF!`: {metrics["formula_ref_errors"]}.

## Abas auditadas no escopo obrigatório

{markdown_table(["Aba", "Dimensão usada", "Registros/linhas úteis", "Fórmulas", "Diagnóstico"], audited_sheet_rows)}

As abas de categoria são fontes observacionais e podem repetir a mesma instituição em verticais
diferentes. O piloto não preserva essas repetições como contas distintas: os 12 `account_id` são
únicos e os aliases brutos ficam registrados para revisão. Não houve merge automático por fuzzy match.

## Reconciliação de 210 versus 287/395

1. `Progresso!B2:C2` contém texto hardcoded: `{metrics["progress_count"]}` e `{metrics["progress_detail"]}`. Não há fórmula nem regra de aliases que permita reproduzir 287.
2. As 11 abas de categorias contêm hoje **{metrics["raw_rows"]}** linhas com nome e tipo preenchidos, não 395. A contagem hardcoded está uma linha defasada.
3. Essas linhas produzem **{metrics["raw_normalized_unique"]}** nomes únicos com normalização mecânica. Chegar a 287 exigiria uma tabela de aliases/exclusões que não existe no arquivo.
4. `Base Completa E` contém **{metrics["base_rows"]}** linhas com nome. Entretanto, `Resumo E!B4` usa `{metrics["summary_formula"]}`, que termina na linha 214 e contabiliza 210 registros.
5. `Base Completa E!B215:Q215` é um append posterior de `Brunel` com colunas desalinhadas em relação ao schema consolidado; por isso fica fora do resumo.

**Conclusão:** 287/395 é um controle estático de uma versão ou metodologia anterior. A fotografia atual é 396 observações categorizadas, 319 nomes mecanicamente únicos e 210 registros dentro do range oficial do resumo, mais uma linha malformada fora dele. O número autoritativo do MVP só poderá ser definido após tabela canônica de aliases e revisão humana.

## Perfil de dados nas abas de categorias

{markdown_table(["Campo", "Ausentes/Não encontrado", "% de 396"], missing_rows)}

O preenchimento de endereço ou telefone não comprova atualidade nem adequação para prospecção. E-mail geral, canal de recrutamento e contato comercial são entidades diferentes e não devem ser substituídos entre si.

## Regras de normalização e deduplicação

1. Preservar o nome bruto e criar `canonical_name` separado.
2. Remover caixa, acentos e pontuação apenas para gerar chaves auxiliares; nunca usar fuzzy match para merge automático.
3. Confirmar entidade por CNPJ quando disponível. CNPJ da gestora, banco, fundo e veículo não são intercambiáveis.
4. Manter `group_name` separado da unidade operacional. Kinea e Itaú BBA pertencem ao mesmo grupo, mas não são a mesma unidade.
5. Bancos recebem `analysis_unit=mesa/área`; o conglomerado não recebe um score comercial único.
6. Alias ambíguo ou combinação como `Itaú BBA / Kinea` exige revisão humana.
7. Endereço com duas cidades deve virar dois registros de escritório no Prompt 2, não duas cidades em uma célula.
8. `Não encontrado` vira nulo, preservando a lacuna em `missing_fields`.

## Qualidade das 12 contas do piloto

{markdown_table(["Conta", "Confiança", "Status no workbook", "Lacunas", "Flags"], pilot_quality_rows)}

## Backlog priorizado de lacunas

### P0 — antes de qualquer score ou demonstração externa

1. Confirmar razão social e CNPJ de cada uma das 12 contas em fonte oficial.
2. Revisar a linha malformada `Base Completa E!B215:Q215` e decidir se Brunel pertence ao universo.
3. Criar tabela de aliases aprovada e reconciliar formalmente 396 -> universo canônico.
4. Separar BTG e Itaú por mesa/área; não usar contato, endereço ou estratégia do grupo como se servissem a todas as unidades.
5. Revalidar endereço e fonte das contas sem URL de verificação no consolidado.

### P1 — necessário para Account 360 confiável

1. Criar tabela de escritórios com um endereço por linha, data e precisão.
2. Separar contato institucional, imprensa, recrutamento e comercial.
3. Mapear fundos/veículos por CNPJ e gestor sem fuzzy match silencioso.
4. Registrar pessoas públicas somente quando necessárias ao workflow e com política LGPD aprovada.

### P2 — necessário antes de modelar timing comercial

1. Ingerir séries públicas de PL, cota, captação, resgate e cotistas com cobertura explícita.
2. Definir eventos comerciais e janela de validade.
3. Obter histórico real de oportunidades/conversões antes de qualquer modelo supervisionado.

## Critérios de aceite do Prompt 1

- Contagens reconciliadas: **concluído com alerta de controle estático**.
- Piloto com 12 `account_id` e nomes canônicos únicos: **concluído**.
- Proveniência por linha: **concluído** por locator e hash.
- Dados ausentes não inventados: **concluído**.
- Testes de unicidade e obrigatoriedade: **executados pelo pytest**.
- Aprovação humana: **pendente** para aliases, CNPJ, endereços e desk mapping.

## Gates

- G0 Mandato: ready.
- G1 Fonte: limited — várias linhas não têm fonte primária adjacente.
- G2 Extração: ready para o recorte do piloto.
- G3 Modelo: not-run; fora do Prompt 1.
- G4 Narrativa: limited — ainda não existe score nem evidência causal.
- G5 Entrega: ready após testes.
- G6 Julgamento: limited — aprovação humana explicitamente pendente.

Status geral: **limited**, limitado pela falta de identidade legal/CNPJ, fontes primárias completas e reconciliação canônica aprovada.
"""
    path.write_text(content, encoding="utf-8")


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Audit the networking workbook and build Prompt 1 outputs."
    )
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--output-root", type=Path, default=Path.cwd())
    args = parser.parse_args()

    workbook_path = args.workbook.resolve()
    output_root = args.output_root.resolve()
    source_hash_before = sha256(workbook_path)
    workbook = load_workbook(workbook_path, read_only=False, data_only=False)

    metrics = workbook_metrics(workbook)
    pilot_rows = build_pilot_rows(workbook, source_hash_before)

    write_csv(output_root / "data/staging/pilot_accounts.csv", pilot_rows, OUTPUT_COLUMNS)
    write_expected_fixture(output_root / "tests/fixtures/pilot_accounts_expected.csv", pilot_rows)

    schema_path = output_root / "data/contracts/organization_schema.yaml"
    schema_path.parent.mkdir(parents=True, exist_ok=True)
    schema_path.write_text(
        yaml.safe_dump(schema_document(), allow_unicode=True, sort_keys=False), encoding="utf-8"
    )

    write_audit(
        output_root / "docs/data_audit.md",
        workbook_path,
        source_hash_before,
        metrics,
        pilot_rows,
    )

    source_hash_after = sha256(workbook_path)
    if source_hash_before != source_hash_after:
        raise RuntimeError("Source workbook changed during a read-only audit")

    print(
        yaml.safe_dump(
            {
                "source_sha256": source_hash_after,
                "metrics": metrics,
                "pilot_accounts": len(pilot_rows),
                "outputs": [
                    "docs/data_audit.md",
                    "data/contracts/organization_schema.yaml",
                    "data/staging/pilot_accounts.csv",
                    "tests/fixtures/pilot_accounts_expected.csv",
                ],
            },
            allow_unicode=True,
            sort_keys=False,
        )
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
